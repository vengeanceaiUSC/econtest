#!/usr/bin/env python3
"""Create Study_Tracker.xlsx — exam topic checklist with green checkoffs."""

import os
import subprocess

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from exam_config import EXAM_STRUCTURE
from excel_pdf_embed import add_pdf_pages_sheet, package_pdf_in_xlsx

OUT = "/workspace/Study_Tracker.xlsx"
PDF_PATH = "/workspace/Labor_Leisure_Problems.pdf"
DIV_PDF_PATH = "/workspace/Diversification_Variance_Problems.pdf"
MU_PDF_PATH = "/workspace/Marginal_Utility_Guide.pdf"
PREF_PDF_PATH = "/workspace/Assumptions_Preferences_Problems.pdf"

PDF_URL = "https://github.com/vengeanceaiUSC/econtest/releases/download/exam-prep-download/Labor_Leisure_Problems.pdf"
LL_XLSX_URL = "https://github.com/vengeanceaiUSC/econtest/releases/download/exam-prep-download/Labor_Leisure_Problems.xlsx"
DIV_PDF_URL = "https://github.com/vengeanceaiUSC/econtest/releases/download/exam-prep-download/Diversification_Variance_Problems.pdf"
DIV_XLSX_URL = "https://github.com/vengeanceaiUSC/econtest/releases/download/exam-prep-download/Diversification_Variance_Problems.xlsx"
MU_PDF_URL = "https://github.com/vengeanceaiUSC/econtest/releases/download/exam-prep-download/Marginal_Utility_Guide.pdf"
MU_XLSX_URL = "https://github.com/vengeanceaiUSC/econtest/releases/download/exam-prep-download/Marginal_Utility_Guide.xlsx"
PREF_PDF_URL = "https://github.com/vengeanceaiUSC/econtest/releases/download/exam-prep-download/Assumptions_Preferences_Problems.pdf"
PREF_XLSX_URL = "https://github.com/vengeanceaiUSC/econtest/releases/download/exam-prep-download/Assumptions_Preferences_Problems.xlsx"

SKIP_GREEN = {"Labor-Leisure", "Diversification"}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
CH4_FILL = PatternFill("solid", fgColor="D6E4F0")
CH9_FILL = PatternFill("solid", fgColor="E2EFDA")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
BOLD = Font(bold=True)
CHECK_FONT = Font(bold=True, size=14, color="006100")
LINK_FONT = Font(color="0563C1", underline="single", size=10)
WRAP = Alignment(wrap_text=True, vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="AAAAAA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "Topic Checklist"

ws["A1"] = "ECON 351 — Exam Topic Tracker"
ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A1:F1")

headers = ["Chapter", "Category", "Questions on Test", "Status", "Notes", "Practice Materials"]
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=3, column=c, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = CENTER
    cell.border = BORDER

row = 4
ch4_total = ch9_total = 0
for section in EXAM_STRUCTURE:
    ch = section["ch"]
    cat = section["category"]
    count = section["count"]
    if ch == 4:
        ch4_total += count
    else:
        ch9_total += count

    done = cat not in SKIP_GREEN
    status = "✓ Done" if done else "— Not yet"
    notes = "" if done else "Still need to practice"
    materials = ""
    if cat == "Labor-Leisure":
        notes = "PDF included in this file (Labor-Leisure PDF tab)"
        materials = "Labor_Leisure_Problems.pdf (embedded)"
    elif cat == "Diversification":
        notes = "PDF included in this file (Diversification PDF tab)"
        materials = "Diversification_Variance_Problems.pdf (embedded)"
    elif cat == "Risk Attitudes":
        notes = "PDF included in this file (Marginal Utility PDF tab)"
        materials = "Marginal_Utility_Guide.pdf (embedded)"
    elif cat == "Assumptions / Preferences":
        notes = "PDF included in this file (Preferences PDF tab)"
        materials = "Assumptions_Preferences_Problems.pdf (embedded)"

    base_fill = CH4_FILL if ch == 4 else CH9_FILL
    row_fill = GREEN_FILL if done else base_fill

    values = [f"Ch. {ch}", cat, count, status, notes, materials]
    for c, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=val)
        cell.fill = row_fill
        cell.border = BORDER
        cell.alignment = WRAP if c != 4 else CENTER
        if c == 4 and done:
            cell.font = CHECK_FONT

    if cat == "Labor-Leisure":
        pdf_cell = ws.cell(row=row, column=6)
        pdf_cell.hyperlink = "#'Labor-Leisure PDF'!A1"
        pdf_cell.value = "Go to embedded PDF tab"
        pdf_cell.font = LINK_FONT
        pdf_cell.fill = row_fill
        pdf_cell.border = BORDER
        note_cell = ws.cell(row=row, column=5)
        note_cell.hyperlink = LL_XLSX_URL
        note_cell.value = "Also: full LL workbook (online)"
        note_cell.font = LINK_FONT
    elif cat == "Diversification":
        pdf_cell = ws.cell(row=row, column=6)
        pdf_cell.hyperlink = "#'Diversification PDF'!A1"
        pdf_cell.value = "Go to embedded PDF tab"
        pdf_cell.font = LINK_FONT
        pdf_cell.fill = row_fill
        pdf_cell.border = BORDER
        note_cell = ws.cell(row=row, column=5)
        note_cell.hyperlink = DIV_XLSX_URL
        note_cell.value = "Also: full Div/Var workbook (online)"
        note_cell.font = LINK_FONT
    elif cat == "Risk Attitudes":
        pdf_cell = ws.cell(row=row, column=6)
        pdf_cell.hyperlink = "#'Marginal Utility PDF'!A1"
        pdf_cell.value = "Go to embedded PDF tab"
        pdf_cell.font = LINK_FONT
        pdf_cell.fill = row_fill
        pdf_cell.border = BORDER
        note_cell = ws.cell(row=row, column=5)
        note_cell.hyperlink = MU_XLSX_URL
        note_cell.value = "Also: full MU guide workbook (online)"
        note_cell.font = LINK_FONT
    elif cat == "Assumptions / Preferences":
        pdf_cell = ws.cell(row=row, column=6)
        pdf_cell.hyperlink = "#'Preferences PDF'!A1"
        pdf_cell.value = "Open Assumptions & Preferences PDF"
        pdf_cell.font = LINK_FONT
        pdf_cell.fill = row_fill
        pdf_cell.border = BORDER
        note_cell = ws.cell(row=row, column=5)
        note_cell.hyperlink = PREF_XLSX_URL
        note_cell.value = "Also: full Preferences workbook (online)"
        note_cell.font = LINK_FONT

    row += 1

# Totals
ws.cell(row=row, column=2, value="Ch. 4 Total").font = BOLD
ws.cell(row=row, column=3, value=ch4_total).font = BOLD
for c in range(1, 7):
    ws.cell(row=row, column=c).border = BORDER
    ws.cell(row=row, column=c).fill = CH4_FILL
row += 1

ws.cell(row=row, column=2, value="Ch. 9 Total").font = BOLD
ws.cell(row=row, column=3, value=ch9_total).font = BOLD
for c in range(1, 7):
    ws.cell(row=row, column=c).border = BORDER
    ws.cell(row=row, column=c).fill = CH9_FILL
row += 1

ws.cell(row=row, column=2, value="GRAND TOTAL").font = BOLD
ws.cell(row=row, column=3, value=24).font = BOLD
for c in range(1, 7):
    ws.cell(row=row, column=c).border = BORDER
    ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor="FFF2CC")

row += 2
ws.cell(row=row, column=1, value="Legend:").font = BOLD
row += 1
ws.cell(row=row, column=1, value="Green = reviewed / done").fill = GREEN_FILL
row += 1
ws.cell(row=row, column=1, value="Not green = still need practice").fill = CH4_FILL

ws.column_dimensions["A"].width = 10
ws.column_dimensions["B"].width = 32
ws.column_dimensions["C"].width = 16
ws.column_dimensions["D"].width = 12
ws.column_dimensions["E"].width = 28
ws.column_dimensions["F"].width = 28
ws.freeze_panes = "A4"

if not os.path.exists(PDF_PATH):
    subprocess.run(["python3", "/workspace/create_labor_leisure_pdf.py"], check=True)
if not os.path.exists(DIV_PDF_PATH):
    subprocess.run(["python3", "/workspace/create_diversification_pdf.py"], check=True)
if not os.path.exists(MU_PDF_PATH):
    subprocess.run(["python3", "/workspace/create_marginal_utility_pdf.py"], check=True)
if not os.path.exists(PREF_PDF_PATH):
    subprocess.run(["python3", "/workspace/create_preferences_pdf.py"], check=True)

add_pdf_pages_sheet(wb, PDF_PATH, sheet_name="Labor-Leisure PDF")
add_pdf_pages_sheet(wb, DIV_PDF_PATH, sheet_name="Diversification PDF")
add_pdf_pages_sheet(wb, MU_PDF_PATH, sheet_name="Marginal Utility PDF")
add_pdf_pages_sheet(wb, PREF_PDF_PATH, sheet_name="Preferences PDF")
wb["Preferences PDF"]["A1"] = "Assumptions & Preferences Problems (full PDF)"
wb.save(OUT)
package_pdf_in_xlsx(OUT, PDF_PATH, internal_name="Labor_Leisure_Problems.pdf")
package_pdf_in_xlsx(OUT, DIV_PDF_PATH, internal_name="Diversification_Variance_Problems.pdf")
package_pdf_in_xlsx(OUT, MU_PDF_PATH, internal_name="Marginal_Utility_Guide.pdf")
package_pdf_in_xlsx(OUT, PREF_PDF_PATH, internal_name="Assumptions_Preferences_Problems.pdf")
print(f"Saved: {OUT} (with embedded study PDFs)")
