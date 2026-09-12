#!/usr/bin/env python3
"""Generate Exam Prep.pptx from Exam Prep.xlsx."""

from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

EXCEL = "/workspace/Exam Prep.xlsx"
OUT = "/workspace/Exam Prep.pptx"

NAVY = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x2E, 0x79, 0xB5)


def set_title(slide, text, subtitle=None):
    slide.shapes.title.text = text
    if subtitle and len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle


def add_bullets(text_frame, lines, size=16, color=None):
    text_frame.clear()
    for i, line in enumerate(lines):
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        p.text = line
        p.level = 0
        p.font.size = Pt(size)
        if color:
            p.font.color.rgb = color


def add_textbox(slide, left, top, width, height, text, size=14, bold=False, color=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    if color:
        p.font.color.rgb = color
    return box


def question_slide(prs, num, label, chapter, category, source, ref, question, notes=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_textbox(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.4),
                f"Q{num} — {label}", size=22, bold=True, color=NAVY)
    meta = f"Ch {chapter} · {category} · {source} ({ref})"
    add_textbox(slide, Inches(0.5), Inches(0.75), Inches(9), Inches(0.35), meta, size=11, color=GRAY)
    add_textbox(slide, Inches(0.5), Inches(1.2), Inches(9), Inches(4.5), question, size=16)
    if notes:
        add_textbox(slide, Inches(0.5), Inches(5.8), Inches(9), Inches(0.8), f"Note: {notes}", size=10, color=GRAY)
    add_textbox(slide, Inches(0.5), Inches(6.5), Inches(9), Inches(0.4),
                "[Work space below — see Answer Key slides at end]", size=10, color=GRAY)


def main():
    wb = load_workbook(EXCEL)
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    set_title(slide, "ECON 351 — Exam Prep", "24-Question Practice Set · Ch.4 & Ch.9")

    # Study plan
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Study Plan — Topic Breakdown")
    lines = [
        "Chapter 4 (12 questions)",
        "  • 2 Consumer Problems",
        "  • 2 Labor-Leisure",
        "  • 2 Intertemporal Consumption",
        "  • 3 Elasticity / Types of Goods",
        "  • 1 Budget Line",
        "  • 2 Assumptions / Preferences",
        "",
        "Chapter 9 (12 questions)",
        "  • 5 Insurance Problems",
        "  • 4 Risk Attitudes",
        "  • 2 Certainty Equivalent & Risk Premium",
        "  • 1 Diversification",
    ]
    add_bullets(slide.placeholders[1].text_frame, lines, size=18)

    # Sources
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Sources")
    add_bullets(slide.placeholders[1].text_frame, [
        "Homework 1 (Ch.4), Homework 2 (Ch.9)",
        "Sample Questions Ch.4 & Ch.9",
        "Mock Midterm 1",
        "Labor-leisure Q3–4: course-style (not in uploads)",
    ], size=18)

    # Practice questions
    ws = wb["Practice Set"]
    questions = []
    for r in range(2, ws.max_row + 1):
        questions.append({
            "num": ws.cell(r, 1).value,
            "label": ws.cell(r, 2).value,
            "ch": ws.cell(r, 3).value,
            "cat": ws.cell(r, 4).value,
            "source": ws.cell(r, 5).value,
            "ref": ws.cell(r, 6).value,
            "q": ws.cell(r, 7).value,
            "notes": ws.cell(r, 10).value or "",
        })

    # Section divider Ch 4
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, Inches(1), Inches(2.5), Inches(8), Inches(1),
                "Chapter 4 — Practice Questions (Q1–Q12)", size=32, bold=True, color=NAVY)

    for q in questions:
        if q["ch"] == 4:
            question_slide(prs, q["num"], q["label"], q["ch"], q["cat"],
                           q["source"], q["ref"], q["q"], q["notes"] if "not in uploads" in str(q["source"]).lower() else None)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, Inches(1), Inches(2.5), Inches(8), Inches(1),
                "Chapter 9 — Practice Questions (Q13–Q24)", size=32, bold=True, color=NAVY)

    for q in questions:
        if q["ch"] == 9:
            question_slide(prs, q["num"], q["label"], q["ch"], q["cat"],
                           q["source"], q["ref"], q["q"])

    # Formula sheet
    ws_f = wb["Formula Sheet"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.5),
                "Formula Sheet — Chapter 4", size=24, bold=True, color=NAVY)
    ch4_lines = []
    for r in range(2, ws_f.max_row + 1):
        topic = ws_f.cell(r, 1).value
        formula = ws_f.cell(r, 2).value
        if topic and str(topic).startswith("CHAPTER 9"):
            break
        if topic and formula and not str(topic).startswith("CHAPTER"):
            ch4_lines.append(f"{topic}: {formula}")
    add_textbox(slide, Inches(0.5), Inches(0.9), Inches(9), Inches(6),
                "\n".join(ch4_lines), size=11)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.5),
                "Formula Sheet — Chapter 9", size=24, bold=True, color=NAVY)
    ch9_lines = []
    started = False
    for r in range(2, ws_f.max_row + 1):
        topic = ws_f.cell(r, 1).value
        formula = ws_f.cell(r, 2).value
        if topic and str(topic).startswith("CHAPTER 9"):
            started = True
            continue
        if started and topic and formula:
            ch9_lines.append(f"{topic}: {formula}")
    add_textbox(slide, Inches(0.5), Inches(0.9), Inches(9), Inches(6),
                "\n".join(ch9_lines), size=11)

    # Answer key section
    ws_a = wb["Answer Key"]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, Inches(1), Inches(2.5), Inches(8), Inches(1),
                "Answer Key", size=36, bold=True, color=NAVY)

    answers = []
    for r in range(2, ws_a.max_row + 1):
        answers.append((
            ws_a.cell(r, 1).value,
            ws_a.cell(r, 2).value,
            ws_a.cell(r, 4).value,
        ))

    # 6 answers per slide
    for i in range(0, len(answers), 6):
        chunk = answers[i:i + 6]
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_textbox(slide, Inches(0.5), Inches(0.3), Inches(9), Inches(0.4),
                    f"Answer Key (Q{chunk[0][0]}–Q{chunk[-1][0]})", size=20, bold=True, color=NAVY)
        y = 0.85
        for num, label, ans in chunk:
            add_textbox(slide, Inches(0.5), Inches(y), Inches(9), Inches(0.25),
                        f"Q{num} — {label}", size=12, bold=True, color=ACCENT)
            add_textbox(slide, Inches(0.5), Inches(y + 0.28), Inches(9), Inches(0.7),
                        str(ans), size=11)
            y += 1.05

    prs.save(OUT)
    print(f"Saved: {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
