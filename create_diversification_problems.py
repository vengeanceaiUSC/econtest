#!/usr/bin/env python3
"""Create Diversification_Variance_Problems.xlsx with embedded PDF."""

import os
import subprocess

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from excel_pdf_embed import add_pdf_pages_sheet, package_pdf_in_xlsx

OUT = "/workspace/Diversification_Variance_Problems.xlsx"
PDF_PATH = "/workspace/Diversification_Variance_Problems.pdf"
PDF_URL = "https://github.com/vengeanceaiUSC/econtest/releases/download/exam-prep-download/Diversification_Variance_Problems.pdf"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
BLUE_FILL = PatternFill("solid", fgColor="D6E4F0")
GREEN_FILL = PatternFill("solid", fgColor="E2EFDA")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
BOLD = Font(bold=True)
LINK_FONT = Font(color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="AAAAAA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

PROBLEMS = [
    {
        "id": "D-1",
        "source": "Sample Ch.9 Q22 | Practice Tests 1 & 3 Q24",
        "bank_id": "Q054",
        "problem": (
            "Feng invests in agricultural companies A & B (rain/drought).\n"
            "Nan invests in airline C & oil company D (oil price up/down).\n\n"
            "a) Are A & B positively or negatively correlated?\n"
            "b) Are C & D positively or negatively correlated?\n"
            "c) Who has the riskier portfolio, Feng or Nan?"
        ),
        "hints": "Positively correlated assets move together -> higher portfolio variance.",
        "answer": "a) Positively correlated. b) Negatively correlated. c) Feng (riskier). Nan always gets $170.",
    },
    {
        "id": "D-2",
        "source": "Mock Midterm Q29 | Practice Test 2 Q24",
        "bank_id": "Q055",
        "problem": (
            "Assets A & B: positively correlated returns.\n"
            "Assets C & D: negatively correlated returns.\n"
            "John invests half in A, half in B.\n"
            "Marc invests half in C, half in D.\n\n"
            "Which is most likely true comparing John and Marc?"
        ),
        "hints": "Negative correlation diversifies -> lower variance for Marc.",
        "answer": "(D) John's investment has higher variance and John is less risk averse than Marc.",
    },
    {
        "id": "D-3",
        "source": "Lecture Ch.9 - Variance formula",
        "bank_id": "Lecture p.6",
        "problem": "Write the variance formula Var(x) and explain what a high variance means.",
        "hints": "Var(x) = sum( pi * (xi - E[x])^2 ). High variance = more risk.",
        "answer": "See lecture slide. Constant variable has Var=0.",
    },
    {
        "id": "D-4",
        "source": "Lecture Ch.9 - Variance examples",
        "bank_id": "Lecture p.7",
        "problem": (
            "Investment 1: 50% $310, 50% $290.\n"
            "Investment 2: 50% $600, 50% $0.\n\n"
            "Compute Var(Investment 1) and Var(Investment 2). Which is riskier?"
        ),
        "hints": "E[x] for Inv1 = 300, Inv2 = 300. Plug into variance formula.",
        "answer": "Var(Inv1)=100. Var(Inv2)=90,000. Investment 2 is much riskier.",
    },
    {
        "id": "D-5",
        "source": "Lecture Ch.9 - Correlation",
        "bank_id": "Lecture p.9",
        "problem": "What is correlation? Give one example of positive correlation.",
        "hints": "Variables that move together. E.g. temperature and ice cream sales.",
        "answer": "Correlation measures co-movement. Correlation does not imply causation.",
    },
    {
        "id": "D-6",
        "source": "Lecture Ch.9 - Diversification",
        "bank_id": "Lecture p.60",
        "problem": (
            "Why does holding both Nvidia (AI exposure) and Sony (gaming/hardware) "
            "reduce risk compared to holding only one stock?"
        ),
        "hints": "Outcomes are not fully correlated - one may do well when the other does poorly.",
        "answer": "Diversification across assets with imperfect correlation lowers portfolio variance.",
    },
]

wb = Workbook()
ws = wb.active
ws.title = "Problems to Solve"
ws["A1"] = "Diversification & Variance Problems - ECON 351"
ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A1:F1")
ws["A2"] = "All problems from sample exam, mock midterm, lecture slides, and practice tests. HW1/HW2 have none."
ws["A2"].font = Font(italic=True, color="666666", size=10)
ws.merge_cells("A2:F2")
ws.cell(row=3, column=1, value="Full PDF (included in this workbook - see 'PDF' tab):").font = BOLD
link = ws.cell(row=3, column=3, value="Open online copy")
link.hyperlink = PDF_URL
link.font = LINK_FONT
ws.merge_cells("A3:B3")
ws.merge_cells("C3:F3")

headers = ["#", "ID", "Source", "Problem", "Hints", "Your work"]
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=5, column=c, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = CENTER
    cell.border = BORDER

for i, p in enumerate(PROBLEMS, 1):
    r = i + 5
    if "Practice Test" in p["source"] or "Sample" in p["source"] or "Mock" in p["source"]:
        fill = YELLOW_FILL
    elif "Lecture" in p["source"]:
        fill = BLUE_FILL
    else:
        fill = GREEN_FILL
    row = [i, p["id"], p["source"], p["problem"], p["hints"], ""]
    for c, val in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.fill = fill
        cell.border = BORDER
        cell.alignment = WRAP

ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 8
ws.column_dimensions["C"].width = 30
ws.column_dimensions["D"].width = 50
ws.column_dimensions["E"].width = 32
ws.column_dimensions["F"].width = 22
ws.freeze_panes = "A6"

ws2 = wb.create_sheet("Answer Key")
ws2["A1"] = "Answer Key"
ws2["A1"].font = Font(bold=True, size=14)
for c, h in enumerate(["ID", "Bank ID", "Answer"], 1):
    cell = ws2.cell(row=3, column=c, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = BORDER
for i, p in enumerate(PROBLEMS, 4):
    for c, val in enumerate([p["id"], p.get("bank_id", ""), p["answer"]], 1):
        cell = ws2.cell(row=i, column=c, value=val)
        cell.border = BORDER
        cell.alignment = WRAP
ws2.column_dimensions["A"].width = 8
ws2.column_dimensions["B"].width = 14
ws2.column_dimensions["C"].width = 65

ws3 = wb.create_sheet("Formula Sheet")
formulas = [
    ("Expected value", "E[x] = sum( pi * xi )"),
    ("Variance", "Var(x) = sum( pi * (xi - E[x])^2 )"),
    ("Risk", "Higher variance = more volatile / riskier"),
    ("Correlation", "Positive: move together | Negative: move opposite"),
    ("Diversification", "Combine assets not fully positively correlated to lower variance"),
    ("Portfolio insight", "Feng (pos. corr.) riskier than Nan (neg. corr.)"),
    ("Mock Q29", "John (pos. corr.) has higher variance than Marc (neg. corr.)"),
]
ws3["A1"] = "Variance, Correlation & Diversification"
ws3["A1"].font = Font(bold=True, size=14)
for i, (topic, formula) in enumerate(formulas, 3):
    ws3.cell(row=i, column=1, value=topic).font = BOLD
    ws3.cell(row=i, column=2, value=formula).alignment = WRAP
ws3.column_dimensions["A"].width = 18
ws3.column_dimensions["B"].width = 55

if not os.path.exists(PDF_PATH):
    subprocess.run(["python3", "/workspace/create_diversification_pdf.py"], check=True)

add_pdf_pages_sheet(wb, PDF_PATH, sheet_name="PDF", first=True)
wb.save(OUT)
package_pdf_in_xlsx(OUT, PDF_PATH, internal_name="Diversification_Variance_Problems.pdf")
print(f"Saved: {OUT} ({len(PROBLEMS)} problems + embedded PDF)")
