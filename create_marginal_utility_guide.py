#!/usr/bin/env python3
"""Create Marginal_Utility_Guide.xlsx with embedded PDF."""

import os
import subprocess

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from excel_pdf_embed import add_pdf_pages_sheet, package_pdf_in_xlsx

OUT = "/workspace/Marginal_Utility_Guide.xlsx"
PDF_PATH = "/workspace/Marginal_Utility_Guide.pdf"
PDF_URL = "https://github.com/vengeanceaiUSC/econtest/releases/download/exam-prep-download/Marginal_Utility_Guide.pdf"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
BLUE_FILL = PatternFill("solid", fgColor="D6E4F0")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
BOLD = Font(bold=True)
LINK_FONT = Font(color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="AAAAAA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

SECTIONS = [
    ("Total Utility (TU)", "Overall satisfaction from a given amount of wealth or goods."),
    ("Marginal Utility (MU)", "Extra satisfaction from exactly one more unit. MU = u'(x)."),
    ("Second derivative", "Measures whether MU is increasing, decreasing, or flat."),
    ("MU increasing example", "MU at 1 = 10, MU at 2 = 20 -> MU rising -> u'' > 0 -> risk-loving"),
    ("NEGATIVE u''", "Concave utility. MU going down. Second derivative negative because MU is decreasing. Risk averse."),
    ("POSITIVE u''", "Convex utility. MU going up, faster every unit. Risk loving."),
    ("NEUTRAL u'' = 0", "Linear utility. Flat MU line. Second derivative is 0. Risk neutral."),
    ("Graphs in PDF", "sqrt(x) concave + decreasing MU | x linear + flat MU | x^2 convex + increasing MU"),
]

wb = Workbook()
ws = wb.active
ws.title = "Summary"
ws["A1"] = "Marginal Utility Guide - ECON 351"
ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A1:C1")
ws.cell(row=2, column=1, value="Full PDF included in 'PDF' tab of this workbook.").font = Font(italic=True, color="666666")
link = ws.cell(row=2, column=3, value="Open online copy")
link.hyperlink = PDF_URL
link.font = LINK_FONT

for c, h in enumerate(["Topic", "Definition"], 1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = BORDER

for i, (topic, definition) in enumerate(SECTIONS, 5):
    ws.cell(row=i, column=1, value=topic).font = BOLD
    ws.cell(row=i, column=1).fill = BLUE_FILL
    ws.cell(row=i, column=1).border = BORDER
    cell = ws.cell(row=i, column=2, value=definition)
    cell.alignment = WRAP
    cell.border = BORDER

ws.column_dimensions["A"].width = 22
ws.column_dimensions["B"].width = 70

if not os.path.exists(PDF_PATH):
    subprocess.run(["python3", "/workspace/create_marginal_utility_pdf.py"], check=True)

add_pdf_pages_sheet(wb, PDF_PATH, sheet_name="PDF", first=True)
# Fix header text for this workbook
wb["PDF"]["A1"] = "Full Marginal Utility Guide PDF"
wb.save(OUT)
package_pdf_in_xlsx(OUT, PDF_PATH, internal_name="Marginal_Utility_Guide.pdf")
print(f"Saved: {OUT} (with embedded PDF)")
