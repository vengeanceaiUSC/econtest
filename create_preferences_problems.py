#!/usr/bin/env python3
"""Create Assumptions_Preferences_Problems.xlsx with embedded PDF."""

import os
import subprocess

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from excel_pdf_embed import add_pdf_pages_sheet, package_pdf_in_xlsx

OUT = "/workspace/Assumptions_Preferences_Problems.xlsx"
PDF_PATH = "/workspace/Assumptions_Preferences_Problems.pdf"
PDF_URL = "https://github.com/vengeanceaiUSC/econtest/releases/download/exam-prep-download/Assumptions_Preferences_Problems.pdf"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
BLUE_FILL = PatternFill("solid", fgColor="D6E4F0")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
BOLD = Font(bold=True)
LINK_FONT = Font(color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin = Side(style="thin", color="AAAAAA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

PROBLEMS = [
    ("P-1a", "HW Ex 2.1", "Q001", "Knives Out vs Top Gun - doesn't care. Violation?", "No - indifference OK"),
    ("P-1b", "HW Ex 2.2", "Q002", "Avatar vs Avengers - cannot decide. Violation?", "Yes - completeness"),
    ("P-1c", "HW Ex 2.3", "Q003", "Parasite > Indy > JP but JP > Parasite. Violation?", "Yes - transitivity"),
    ("P-2", "Sample Q5", "Q004", "Slope of IC tells you completeness / MRS / prices / transitivity?", "(C) MRS"),
    ("P-3", "Sample Q6", "Q007", "Magnitude of IC slope equals what?", "(A) MRS"),
    ("P-4", "Sample Q7", "—", "Inka linear IC through (12,14) - Y intercept?", "(B) Y = 17.333"),
    ("P-5", "Sample Q8", "Q005", "Serena IC slope -7/5. Utility function?", "(D) X/5 + Y/7"),
    ("P-6", "Sample Q9", "—", "Angelique Cobb-Douglas same IC - find XB", "(D) XB = 16"),
    ("P-7", "Sample Q30", "Q006", "Tony perfect complements 1/4 X with 1/6 Y. Optimal mix?", "(D) 6X = 4Y"),
    ("P-8", "Sample Q31", "—", "Jimmy L-shaped IC. Optimal mix?", "(A) 9X = 4Y"),
    ("P-9", "Mock Q22", "—", "u = X^6 Y^3 same IC - find XB", "(D) XB = 16"),
]

wb = Workbook()
ws = wb.active
ws.title = "Problems"
ws["A1"] = "Assumptions & Preferences Problems - ECON 351"
ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A1:F1")
ws.cell(row=2, column=1, value="Full PDF included in 'PDF' tab.").font = Font(italic=True, color="666666")
link = ws.cell(row=2, column=3, value="Open online copy")
link.hyperlink = PDF_URL
link.font = LINK_FONT

headers = ["ID", "Source", "Bank ID", "Question summary", "Answer"]
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = BORDER

for i, row in enumerate(PROBLEMS, 5):
    fill = YELLOW_FILL if "HW" in row[1] or "Practice" in row[1] else BLUE_FILL
    for c, val in enumerate(row, 1):
        cell = ws.cell(row=i, column=c, value=val)
        cell.fill = fill
        cell.border = BORDER
        cell.alignment = WRAP

ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 10
ws.column_dimensions["D"].width = 48
ws.column_dimensions["E"].width = 28
ws.freeze_panes = "A5"

if not os.path.exists(PDF_PATH):
    subprocess.run(["python3", "/workspace/create_preferences_pdf.py"], check=True)

add_pdf_pages_sheet(wb, PDF_PATH, sheet_name="PDF", first=True)
wb["PDF"]["A1"] = "Full Assumptions & Preferences PDF"
wb.save(OUT)
package_pdf_in_xlsx(OUT, PDF_PATH, internal_name="Assumptions_Preferences_Problems.pdf")
print(f"Saved: {OUT} (with embedded PDF)")
