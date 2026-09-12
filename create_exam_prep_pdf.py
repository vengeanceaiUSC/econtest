#!/usr/bin/env python3
"""Generate Exam Prep.pdf from Exam Prep.xlsx."""

from openpyxl import load_workbook
from fpdf import FPDF

EXCEL = "/workspace/Exam Prep.xlsx"
OUT = "/opt/cursor/artifacts/Exam_Prep.pdf"


def sanitize(text):
    if text is None:
        return ""
    s = str(text)
    for a, b in [
        ("\u2013", "-"), ("\u2014", "-"), ("\u2018", "'"), ("\u2019", "'"),
        ("\u201c", '"'), ("\u201d", '"'), ("\u2026", "..."), ("\u00b2", "^2"),
        ("\u00b3", "^3"), ("\u221a", "sqrt"), ("\u00b7", "*"), ("\u2192", "->"),
        ("\u2113", "l"), ("\u03b1", "a"), ("\u03b2", "b"),
    ]:
        s = s.replace(a, b)
    return s.encode("latin-1", "replace").decode("latin-1")


class PDF(FPDF):
    def header(self):
        self.set_font("Helvetica", "B", 10)
        self.set_text_color(80, 80, 80)
        self.cell(0, 8, "ECON 351 - Exam Prep (24-Question Practice Set)", align="C", new_x="LMARGIN", new_y="NEXT")
        self.ln(2)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 10, f"Page {self.page_no()}", align="C")

    def section_title(self, title):
        self.set_font("Helvetica", "B", 14)
        self.set_text_color(31, 78, 121)
        self.multi_cell(0, 8, sanitize(title))
        self.ln(2)

    def body_text(self, text, size=11, bold=False):
        self.set_font("Helvetica", "B" if bold else "", size)
        self.set_text_color(30, 30, 30)
        self.multi_cell(0, 5.5, sanitize(text))
        self.ln(1)

    def meta_text(self, text):
        self.set_font("Helvetica", "", 9)
        self.set_text_color(100, 100, 100)
        self.multi_cell(0, 5, sanitize(text))
        self.ln(2)


def main():
    wb = load_workbook(EXCEL)
    pdf = PDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    # Cover
    pdf.set_font("Helvetica", "B", 22)
    pdf.set_text_color(31, 78, 121)
    pdf.ln(30)
    pdf.cell(0, 12, "ECON 351 Exam Prep", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 14)
    pdf.set_text_color(60, 60, 60)
    pdf.cell(0, 10, "24-Question Practice Set", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(10)
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 6, sanitize(
        "Sources: Homework 1 & 2, Sample Questions Ch.4 & Ch.9, Mock Midterm 1.\n"
        "Topic counts match professor breakdown (12 Ch.4 + 12 Ch.9).\n"
        "Labor-leisure Q3-Q4 are course-style (not in uploads)."
    ), align="C")

    # Study plan
    pdf.add_page()
    pdf.section_title("Study Plan")
    plan = [
        ("Chapter 4 (12)", [
            "2 Consumer Problems", "2 Labor-Leisure", "2 Intertemporal Consumption",
            "3 Elasticity / Types of Goods", "1 Budget Line", "2 Assumptions / Preferences",
        ]),
        ("Chapter 9 (12)", [
            "5 Insurance Problems", "4 Risk Attitudes",
            "2 Certainty Equivalent & Risk Premium", "1 Diversification",
        ]),
    ]
    for heading, items in plan:
        pdf.body_text(heading, bold=True)
        for item in items:
            pdf.body_text(f"  - {item}")
        pdf.ln(2)

    # Practice questions
    ws = wb["Practice Set"]
    current_ch = None
    for r in range(2, ws.max_row + 1):
        ch = ws.cell(r, 3).value
        if ch != current_ch:
            pdf.add_page()
            pdf.section_title(f"Chapter {ch} - Practice Questions")
            current_ch = ch
        num = ws.cell(r, 1).value
        label = ws.cell(r, 2).value
        cat = ws.cell(r, 4).value
        source = ws.cell(r, 5).value
        ref = ws.cell(r, 6).value
        question = ws.cell(r, 7).value
        notes = ws.cell(r, 10).value or ""

        pdf.body_text(f"Q{num} - {label}", bold=True)
        pdf.meta_text(f"{cat} | {source} ({ref})")
        pdf.body_text(str(question))
        if notes:
            pdf.meta_text(f"Note: {notes}")
        pdf.ln(4)
        # workspace line
        pdf.set_draw_color(200, 200, 200)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y())
        pdf.ln(6)

    # Formula sheet
    pdf.add_page()
    pdf.section_title("Formula Sheet - Chapter 4")
    ws_f = wb["Formula Sheet"]
    for r in range(2, ws_f.max_row + 1):
        topic = ws_f.cell(r, 1).value
        formula = ws_f.cell(r, 2).value
        note = ws_f.cell(r, 3).value or ""
        if topic and str(topic).startswith("CHAPTER 9"):
            break
        if topic and formula and not str(topic).startswith("CHAPTER"):
            pdf.body_text(f"{topic}: {formula}" + (f"  ({note})" if note else ""), size=10)

    pdf.add_page()
    pdf.section_title("Formula Sheet - Chapter 9")
    started = False
    for r in range(2, ws_f.max_row + 1):
        topic = ws_f.cell(r, 1).value
        formula = ws_f.cell(r, 2).value
        note = ws_f.cell(r, 3).value or ""
        if topic and str(topic).startswith("CHAPTER 9"):
            started = True
            continue
        if started and topic and formula:
            pdf.body_text(f"{topic}: {formula}" + (f"  ({note})" if note else ""), size=10)

    # Answer key
    pdf.add_page()
    pdf.section_title("Answer Key")
    ws_a = wb["Answer Key"]
    for r in range(2, ws_a.max_row + 1):
        num = ws_a.cell(r, 1).value
        label = ws_a.cell(r, 2).value
        ans = ws_a.cell(r, 4).value
        pdf.body_text(f"Q{num} - {label}", bold=True)
        pdf.body_text(str(ans), size=10)
        pdf.ln(2)

    pdf.output(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
