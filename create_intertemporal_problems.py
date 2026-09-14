#!/usr/bin/env python3
"""Create Intertemporal_Consumption_Problems.xlsx with embedded PDF."""

import os
import subprocess

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from excel_pdf_embed import add_pdf_pages_sheet, package_pdf_in_xlsx

OUT = "/workspace/Intertemporal_Consumption_Problems.xlsx"
PDF_PATH = "/workspace/Intertemporal_Consumption_Problems.pdf"
PDF_URL = "https://github.com/vengeanceaiUSC/econtest/releases/download/exam-prep-download/Intertemporal_Consumption_Problems.pdf"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
BLUE_FILL = PatternFill("solid", fgColor="D6E4F0")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
LINK_FONT = Font(color="0563C1", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
thin = Side(style="thin", color="AAAAAA")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

PROBLEMS = [
    ("IT-1a", "HW Ex 5.1", "Q023", "Contract PV comparison (10M+90M vs 42M+54M, r=0.20)", "Contract 1 cheaper (PV=85M)"),
    ("IT-1b", "HW Ex 5.2", "—", "Utility U=2c1^0.4 c2^0.6 - patient? MU ratio?", "Values future more; MUc1/MUc2=(2/3)(c2/c1)"),
    ("IT-1c", "HW Ex 5.3-5.4", "Q024", "Optimal c1*, c2* under contract 1", "c1*=34, c2*=61.2"),
    ("IT-1d", "HW Ex 5.5", "—", "Optimal c1*, c2* under contract 2", "c1*=34.8, c2*=62.64"),
    ("IT-1e", "HW Ex 5.6", "—", "Which contract does Shohei prefer?", "Compare utility from each contract"),
    ("IT-2", "Sample Q41", "Q025", "I1=10, I2=12, r=0.5, u=q1*q2 - save or borrow?", "q1*=9, q2*=13.5; saves 1"),
    ("IT-3", "Sample Q42", "—", "Same as Q41 but r=0.05", "q1*=75/7; borrows 5/7"),
    ("IT-4", "Mock Q24", "Q026", "Intertemporal budget constraint (I1=60, I2=125, r=0.25)", "(A) q1 + 0.8q2 = 160"),
    ("IT-5", "Mock Q25", "Q027", "Optimal q1 from Q24 setup", "(D) q1 = 80"),
]

wb = Workbook()
ws = wb.active
ws.title = "Problems"
ws["A1"] = "Intertemporal Consumption Problems - ECON 351"
ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A1:F1")
ws.cell(row=2, column=1, value="Full PDF included in 'PDF' tab.").font = Font(italic=True, color="666666")
link = ws.cell(row=2, column=3, value="Open online copy")
link.hyperlink = PDF_URL
link.font = LINK_FONT

headers = ["ID", "Source", "Bank ID", "Question summary", "Answer"]
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = BORDER

for i, row in enumerate(PROBLEMS, 5):
    fill = YELLOW_FILL if "HW" in row[1] or "Practice" in row[1] else BLUE_FILL
    for c, val in enumerate(row, 1):
        cell = ws.cell(row=i, column=c, value=val)
        cell.fill = fill
        cell.border = BORDER
        cell.alignment = WRAP

ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 10
ws.column_dimensions["D"].width = 48
ws.column_dimensions["E"].width = 32
ws.freeze_panes = "A5"

if not os.path.exists(PDF_PATH):
    subprocess.run(["python3", "/workspace/create_intertemporal_pdf.py"], check=True)

add_pdf_pages_sheet(wb, PDF_PATH, sheet_name="PDF", first=True)
wb["PDF"]["A1"] = "Full Intertemporal Consumption PDF"
wb.save(OUT)
package_pdf_in_xlsx(OUT, PDF_PATH, internal_name="Intertemporal_Consumption_Problems.pdf")
print(f"Saved: {OUT} (with embedded PDF)")
