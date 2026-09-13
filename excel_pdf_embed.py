#!/usr/bin/env python3
"""Embed a PDF into an Excel workbook as page images + packaged file."""

import os
import shutil
import zipfile

import pymupdf
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font


def add_pdf_pages_sheet(wb, pdf_path, sheet_name="PDF (full document)", first=False):
    """Render each PDF page as an image on a dedicated worksheet."""
    if not os.path.exists(pdf_path):
        return False

    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.delete_rows(1, ws.max_row)
    elif first:
        ws = wb.create_sheet(sheet_name, 0)
    else:
        ws = wb.create_sheet(sheet_name)

    ws.sheet_view.showGridLines = False
    ws["A1"] = "Full PDF below (same as Labor_Leisure_Problems.pdf)"
    ws["A1"].font = Font(bold=True, size=12)
    row_px = 30  # pixels from top for first page

    cache_dir = os.path.join(os.path.dirname(pdf_path), ".pdf_page_cache")
    os.makedirs(cache_dir, exist_ok=True)

    doc = pymupdf.open(pdf_path)
    for i, page in enumerate(doc):
        mat = pymupdf.Matrix(1.4, 1.4)
        pix = page.get_pixmap(matrix=mat, alpha=False)
        img_path = os.path.join(cache_dir, f"{os.path.basename(pdf_path)}_page_{i + 1}.png")
        pix.save(img_path)

        img = XLImage(img_path)
        max_w = 720
        if img.width > max_w:
            ratio = max_w / img.width
            img.width = int(img.width * ratio)
            img.height = int(img.height * ratio)

        anchor_row = max(1, int(row_px / 15))
        ws.add_image(img, f"A{anchor_row}")
        row_px += img.height + 24

    ws.column_dimensions["A"].width = 105
    return True


def package_pdf_in_xlsx(xlsx_path, pdf_path, internal_name="Labor_Leisure_Problems.pdf"):
    """Add the raw PDF bytes inside the .xlsx zip package."""
    if not os.path.exists(pdf_path):
        return False

    tmp_path = f"{xlsx_path}.tmp"
    embed_path = f"xl/embeddings/{internal_name}"
    with zipfile.ZipFile(xlsx_path, "r") as zin:
        with zipfile.ZipFile(tmp_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, zin.read(item.filename))
            zout.write(pdf_path, embed_path)
    shutil.move(tmp_path, xlsx_path)
    return True
